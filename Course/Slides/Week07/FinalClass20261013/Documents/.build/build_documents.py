from pathlib import Path
import hashlib,json,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
ROOT=Path(__file__).resolve().parents[1];BUILD=ROOT/'.build';BUILD.mkdir(exist_ok=True)
SOURCE=ROOT.parents[2]/'Week01/Lab/Lab 1 Answer Key.docx'
# ROOT is .../FinalClass20261013/Documents
SOURCE=Path(r'C:/Users/antho/Perforce/anthony.monteagudo_TechArt/anthony.monteagudo/Course/Slides/Week01/Lab/Lab 1 Answer Key.docx')
REF=BUILD/'reference.docx';shutil.copy2(SOURCE,REF)
CAP=Path(r'C:/Users/antho/Perforce/anthony.monteagudo_TechArt/anthony.monteagudo/Week07UnrealPython/Saved/LiveRehearsal')
BLUE='075A9E';INK='172235';COLORS=['0A2D62','075A9E','168A86','397B35','C66C0A'];PALE=['E8EEF7','E7F0FA','E4F4F2','EAF4E8','FFF1DB']
(BUILD/'artifact.md').write_text(f'''# Week 7 handout design contract
Reference: retained reference.docx; original {SOURCE}
SHA256: {hashlib.sha256(REF.read_bytes()).hexdigest()}
Authority: the current eight-page Week 1 Lab answer key; user explicitly requested its standard.
Landscape Letter 11 x 8.5 inches; source margins retained (approximately 0.28 top, 0.24 bottom, 0.42 sides). Arial; navy text; blue title band; five colored reasoning stages; pale matrix rows; concise contextual framing and substantive examples. Preserve styles, theme, section geometry and header/footer parts by cloning the retained DOCX. Replace body, source-specific images and prose; no Week 1 scenario answers reused. Intentional adaptations: larger student writing areas; standalone task/review/synthesis sheets; actual current Unreal screenshots in instructor run sheet. Title band follows reference rather than generic black title preset. Render with Word COM after packaged renderer failed because soffice.exe is unavailable. Inspect every final page PNG.
''',encoding='utf-8')
def shade(cell,color):
 p=cell._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),color);p.append(s)
def text(cell,s,bold=False,color=INK,size=10):
 cell.text='';p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
 r=p.add_run(s);r.bold=bold;r.font.name='Arial';r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color)
 return p

def new():
 d=Document(REF)
 for child in list(d._element.body):
  if child.tag!=qn('w:sectPr'):d._element.body.remove(child)
 d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
 d.styles['Normal'].paragraph_format.space_after=Pt(5)
 for sec in d.sections:
  for p in sec.footer.paragraphs:p.text='Week 7  |  Technical Art  |  Classroom working materials'
 d.core_properties.title='Week 7 Technical Art';d.core_properties.author='Prof. Monteagudo';d.core_properties.comments=''
 return d

def p(d,s='',bold=False,size=10):
 v=d.add_paragraph();v.paragraph_format.space_after=Pt(5);v.paragraph_format.space_before=Pt(0)
 r=v.add_run(s);r.bold=bold;r.font.name='Arial';r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(INK);return v

def page(d,label,title,intro):
 if len(d.paragraphs) or len(d.tables):d.add_page_break()
 p(d,label.upper(),True,8)
 t=d.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Inches(10.12);shade(t.cell(0,0),BLUE);text(t.cell(0,0),title,True,'FFFFFF',19)
 p(d,intro,size=10)

def table(d,headers,rows,widths=None,heights=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,h in enumerate(headers):shade(t.cell(0,i),COLORS[0]);text(t.cell(0,i),h,True,'FFFFFF',9)
 for j,row in enumerate(rows):
  rr=t.add_row()
  for i,v in enumerate(row):shade(rr.cells[i],PALE[j%len(PALE)]);text(rr.cells[i],v,size=9)
  if heights:rr.height=Inches(heights[j] if isinstance(heights,list) else heights)
 for rr in t.rows:
  pr=rr._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
  for cell in rr.cells:cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 p(d,'',size=2);return t

def matrix(d,rows,write=False):
 t=d.add_table(rows=0,cols=2);t.autofit=False;t.columns[0].width=Inches(1.45);t.columns[1].width=Inches(8.67)
 for i,(label,body) in enumerate(rows):
  rr=t.add_row();shade(rr.cells[0],COLORS[i%5]);shade(rr.cells[1],PALE[i%5]);text(rr.cells[0],label,True,'FFFFFF',10);text(rr.cells[1],body,size=9)
  rr.height=Inches(.48 if write else .38)
  pr=rr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
 p(d,'',size=2)

def save(d,name):
 for t in d.tables:
  widths=[c.width for c in t.columns]
  for row in t.rows:
   for i,cell in enumerate(row.cells):
    cell.width=widths[i]
    pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for side,val in [('top',45),('bottom',45),('left',90),('right',90)]:
     e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');m.append(e)
    pr.append(m)
 for sec in d.sections:
  for par in sec.footer.paragraphs:
   for run in par.runs:run.font.size=Pt(7)
 d.save(ROOT/name)

# STUDENT LAB 1
D=new();page(D,'Lab 1 student handout   8 minutes   fictional case','Investigate a batch renamer','Decide what the tool should do before seeing its output. Then use evidence to identify one issue and design a retest. No coding or installation is needed.')
p(D,'Group ______   Reader ______   Recorder ______   Evidence checker ______   Reporter ______')
matrix(D,[('Task and scope','The button says Rename selected. Give only selected assets exactly one SM_ prefix. Preserve the rest of each name. Already-prefixed names and all unselected assets must stay unchanged.')])
table(D,['Selected','Starting name','Expected before reveal','Observed after reveal'],[['Yes','crate_A','',''],['Yes','crate_B','',''],['Yes','SM_crate_C','',''],['No','barrel_D','','']],[.8,2.3,3.5,3.5],.35)
matrix(D,[('1  Problem','Before the reveal: what must stay unchanged, and what evidence would make you trust the result?\n________________________________________________________________________________'),('2  Investigate','After the reveal: With [input], I expected [result], but observed [result].\n________________________________________________________________________________'),('3  Intervention','What would you change first, and why? Separate the observed behavior from a guessed cause.\n________________________________________________________________________________'),('4  Impact','Retest: starting input + action + exact expected result, including the unselected asset.\n________________________________________________________________________________'),('5  Tradeoff','What remains unknown? Consider a second run, Undo or empty selection. What would testing it cost?\n________________________________________________________________________________')],True)
p(D,'Prepare one 20-second finding. Keep your first judgment; revisit it during the UX discussion.',True,9);save(D,'Week 7 Lab 1 Student Handout.docx')

# PEER PACKET
D=new();page(D,'Lab 2 preparation   one per author','Prepare one task for peer testing','Give every reviewer the same small task. A goal, safe starting data and observable success make three reviews comparable. Prepare for four minutes before the briefing.')
p(D,'Author __________________  Tool and version __________________  Group ____  Letter A / B / C / D')
matrix(D,[('1  Problem','Who uses this tool, and what should they achieve? State a goal without teaching the clicks.\n________________________________________________________________________________\n________________________________________________________________________________'),('2  Investigate','Starting data, machine and prerequisites. What must remain unchanged?\n________________________________________________________________________________\n________________________________________________________________________________'),('3  Intervention','The exact task you will give the reviewer, in one or two sentences.\n________________________________________________________________________________\n________________________________________________________________________________'),('4  Impact','What visible evidence counts as success? How will you check the result?\n________________________________________________________________________________\n________________________________________________________________________________'),('5  Tradeoff','Reset steps and fallback. Declare hands-on, recorded demonstration or paper/concept review.\n________________________________________________________________________________\n________________________________________________________________________________')],True)
p(D,'Example: “Create a material instance from this texture set and assign it to the target cube. The protected cube must stay unchanged.” Success includes the visible result and correct scope, not just a success message.',size=10)
p(D,'Use prepared copies. If the tool cannot run, show recorded evidence or a concept. Label execution as untested when the reviewer cannot execute it. Do not install software during the round.',True,9)
page(D,'Lab 2 review sheet   print three copies per student','Test a peer tool','Reviewer: predict, try the task and record what happens. Author: observe before explaining; give needed help and make that help visible in the record. Give this sheet to the author.')
p(D,'Round 1 / 2 / 3   Group ____   Author __________________   Reviewer __________________')
p(D,'Tool and version ____________________   Mode: hands-on / recording / paper or concept')
matrix(D,[('Task','Goal, starting input and success evidence. What must stay unchanged?\n________________________________________________________________________________\n________________________________________________________________________________'),('Observation','Action + expected result + observed result. Label recorded or imagined responses clearly.\n________________________________________________________________________________\n________________________________________________________________________________'),('Help and result','Result: independent / with help / blocked / execution not tested. Where did you hesitate?\n________________________________________________________________________________\n________________________________________________________________________________'),('Keep and priority','One evidenced strength. One priority issue or open question and why it matters. If no issue was observed, say so.\n________________________________________________________________________________\n________________________________________________________________________________'),('Next test','Starting input + action + expected evidence. Proposed or attempted? If attempted, record the result.\n________________________________________________________________________________\n________________________________________________________________________________')],True)
p(D,'Each direction takes four minutes: 0:00 task; 0:30 attempt; 2:30 record; 3:30 explain and hand over. Swap at 4:00. Finish and reset at 8:00. Describe behavior, not the person. This is an observation record, not a numeric grade.',True,9)
page(D,'Lab 2 synthesis   final five minutes','Choose a revision and a retest','Read your reviews, compare evidence and choose one next step. Repeated feedback is useful evidence; a single serious scope or recovery failure can also deserve priority.')
p(D,'Author __________________   Distinct peer reviewers ____   Review modes __________________')
table(D,['Review','Evidence or uncertainty','Consequence'],[['1','',''],['2','',''],['3 or unavailable','','']],[1.3,5.2,3.6],.50)
matrix(D,[('Priority','Which issue or uncertainty matters most, and why?\n________________________________________________________________________________'),('Revision','What specific change will you make or investigate?\n________________________________________________________________________________'),('Retest','State input + action + expected result and protected scope. What would disprove your fix?\n________________________________________________________________________________'),('Changed judgment','What did peer evidence change about your first judgment of the tool?\n________________________________________________________________________________'),('Remaining question','What is still untested? Name the next evidence needed.\n________________________________________________________________________________')],True)
p(D,'A recording can show a result; a concept can expose misunderstandings. Neither establishes successful hands-on execution by the reviewer. Existing course policy governs submission; this packet creates no new deadline or grade.',size=9)
save(D,'Week 7 Peer Review Packet.docx')

# FACILITATOR
D=new();page(D,'Instructor guide   1 of 2','Facilitate Lab 1','Use the fictional renamer to establish evidence-based review before teaching UX vocabulary. Students compare a stated contract with visible behavior, then propose a test. Accept different priorities when the reasoning is supported.')
table(D,['Minute','Facilitation'],[['0–1','Distribute one group sheet. Assign reader, recorder, evidence checker and reporter. Read the rule.'],['1–2','Individual prediction. Do not reveal results.'],['2–4','Compare expected names as a group. Check that protected scope is explicit.'],['4–5','Record what evidence would establish success.'],['5–6','Open reveal.html. Display the fictional “Success: 4 assets updated” message and results.'],['6–7','Write one supported finding and a retest.'],['7–8','Two 20-second reports, then 20 seconds for questions.']],[1,9.1])
table(D,['Input','Expected','Observed'],[['crate_A selected','SM_crate_A','SM_crate_A'],['crate_B selected','SM_crate_B','SM_crate_B'],['SM_crate_C selected','SM_crate_C','SM_SM_crate_C'],['barrel_D unselected','barrel_D','SM_barrel_D']],[3.2,3.4,3.5])
matrix(D,[('Interpretation','The output violates both protected scope and the exactly-one-prefix rule. The success count alone proves neither. Internal cause, Undo, a second run and performance remain unknown.'),('Model finding','With only the three crates selected, I expected barrel_D to stay unchanged, but it became SM_barrel_D. This changes unrelated work.'),('Model retest','Restore all four starting names; select only the crates; run once; inspect all four expected names above. Then propose a separate second-run test. Do not claim that test has occurred.'),('Acceptable alternatives','Prioritize scope because unrelated work changes, or the duplicate prefix because it violates the naming contract. Ask for the consequence and supporting observation, not a preferred phrase.'),('Discussion bridge','Which preview evidence could expose the mistake before Apply? What does a success message omit? Revisit first judgments during UX, then carry input/action/evidence/protected scope into peer review.')])
page(D,'Instructor guide   2 of 2','Facilitate reciprocal peer reviews','Protect task preparation and complete attempts. Form groups of four, assign A–D, and keep tools on their prepared machines. Each pair tests in both directions. The later letter tests the earlier letter’s tool first.')
table(D,['Time','Pairing or action'],[['5:16–5:20','Prepare one bounded task and a fallback.'],['5:20–5:23','Brief roles and modes, confirm letters and answer questions.'],['5:23–5:33','Round 1: A–B and C–D.'],['5:33–5:34','Switch partners.'],['5:34–5:44','Round 2: A–C and B–D.'],['5:44–5:45','Switch partners.'],['5:45–5:55','Round 3: A–D and B–C.'],['5:55–6:00','Read reviews; choose priority, revision and retest; record changed judgment.']],[1.5,8.6])
p(D,'Call every round: 0:00 state task; 0:30 attempt; 2:30 record; 3:30 explain; 4:00 swap; 4:30 attempt; 6:30 record; 7:30 explain; 8:00 finish and reset; 10:00 switch.',True)
matrix(D,[('Setup','Per student: preparation page, three copies of review page, synthesis page. One Lab 1 sheet per group. Display dashboard.html. Example for 24: six groups, 72 review sheets; roster is not assumed.'),('Coach evidence','“What did you expect, what happened, and why does it matter?” Record author help. A strength needs evidence too. If no fault is observed, propose a test rather than invent a defect.'),('Handle blockers','Allow at most one minute of troubleshooting within the attempt, then declare blocked and switch to recording or concept. Stop risks to real work. No software installation during rounds.'),('Attendance','Trio: A–B, A–C, B–C; third observes, each gets two distinct reviews. Five: A–B/C–D; A–C/B–E; A–D/C–E. A/C get three; B/D/E get two. Do not count observers or repeated partners as new reviewers.'),('Protect the close','End UX at 5:16. If late, state the reduced number of complete rounds honestly; retain five minutes of synthesis. Full staff scripts and additional attendance cases are in facilitator.html.')])
save(D,'Week 7 Lab Instructor Guide.docx')

# LIVE RUN SHEET
D=new();page(D,'Live Unreal run sheet   1 of 4   setup','Prepare the live Material and VFX demo','Teach the current Material & VFX Tools in Unreal 5.7.4. The native Slate tab and PySide window share the same Python planner and live backend. Use the isolated classroom fixture and preserve the saved baseline.')
matrix(D,[('Before class','Launch Scripts/launch_live_demo.ps1 from Week07UnrealPython. It opens W07_MaterialVFXLab and runs the non-destructive preflight. If this project is already open, use its existing editor and Tools > Material & VFX Tools.'),('Check the fixture','Two plain cubes, three empty pedestals, no placed effects. The near cube is W07_Crate_Target; the far cube is W07_Crate_Protected. Both use the same shared mesh. Open the map from Content/Week07/Maps if needed.'),('Prepare inputs','Select W07_Crate_Target in the Outliner. In Content/Week07/LiveDemo/Textures select all three T_Crate textures: BaseColor, Normal and ORM. Content Browser selection and level selection are separate. Wait for browser selection to update before Preview.'),('Panel and naming','Tools > Material & VFX Tools. Material Setup, slot 0, fresh MI_ name such as MI_Crate_Live_01. Do not overwrite an existing instance. For PySide: Python console: import week07_tools as w7; w7.show_assistants().'),('Fallback readiness','Open live-demo.html in a browser before teaching. It contains actual baseline, material, VFX and native-panel captures, plus read-aloud prompts. It works offline without Unreal or Qt. If a failure persists beyond 30 seconds, use it within the same demo slot.')])
p(D,'Optional selection helpers in the Unreal Python console',True,12)
p(D,'from week07_tooling.editor import live_demo\nprint(live_demo.preflight())\nlive_demo.select_material()\nlive_demo.select_vfx()',size=10)
p(D,'Run one helper at a time; Preview only after the browser updates. Helpers select inputs and frame the camera. They do not apply changes, reset the map, delete assets or save.',size=10)
p(D,'Go/no-go: baseline visible, both workflow previews available, fallback already open. Native plugin binaries match UE 5.7.4 on this Windows machine; another engine build requires rebuilding the plugin. Use the offline fallback if that prerequisite fails.',True,10)
page(D,'Live Unreal run sheet   2 of 4   material   3:49–3:56','Demonstrate material setup','Ask students to predict scope before Apply. The goal is one new material instance and one component override, while the second cube and shared mesh stay unchanged.')
table(D,['Elapsed','Do and say'],[['0:00–1:00','Show both cubes. “They share a mesh. Which one should change, and what must remain untouched?” Point to the selected target and three textures.'],['1:00–2:30','Choose Material Setup, slot 0 and a fresh MI_ name. Preview. Read the destination, parent, three texture mappings and one target. “Preview has not changed the scene.”'],['2:30–4:00','Apply once. Show the near blue cube and far plain cube. Check the instance parent and texture parameters if time permits. “A success message needs a visible scope check.”'],['4:00–5:00','Try Preview again with the same output name. It must block overwrite. “What should happen on a repeat?”'],['5:00–6:00','Use Edit > Undo only when its label identifies the material assignment. Target returns to plain. The new instance remains. Explain that assignment Undo and asset cleanup are separate.'],['6:00–7:00','Leave the unsaved draft for later inspection; use a fresh name next time. Ask one student to state input, action, expected result and protected scope.']],[1.4,8.7])
D.add_picture(str(CAP/'02_material.png'),width=Inches(4.3));p(D,'Actual Unreal capture: target becomes blue; protected cube remains plain. The scene was inspected in the running editor, not simulated for this image.',size=9)
p(D,'3:56–3:58 UI comparison: native Slate belongs to the editor tab system; PySide offers the current staged interface. Both call the same workflow. Use existing Stage 1/2/3 practice captures for the brief visual comparison; do not repeat the whole operation.',size=10)
page(D,'Live Unreal run sheet   3 of 4   VFX   3:58–4:02','Demonstrate bounded VFX placement','Place the prepared classroom fountain at three tagged anchors. Separate predictable placement from runtime performance: a bounded batch does not establish a frame-time budget.')
table(D,['Elapsed','Do and say'],[['0:00–0:45','Select W07_VFXAnchor_1, _2 and _3. Choose VFX Placement, classroom fountain, scale 1.0. “How many effects should this create, and where?”'],['0:45–1:30','Preview: three anchors, three effects, their positions and rotations. Apply once. Show three visible fountains.'],['1:30–2:15','Reselect the same anchors if necessary. Preview again: occupied anchors must block repeated placement. “The tool checks duplication before mutation.”'],['2:15–3:00','Edit > Undo the known VFX placement. Verify all three effects disappear; pedestals remain. If Undo label is unfamiliar, stop and use recovery instructions.'],['3:00–4:00','Ask: “What have we verified? What remains unmeasured?” Answer: placement and recovery checked; runtime performance on target hardware remains unmeasured. Hand off at 4:02.']],[1.4,8.7])
D.add_picture(str(CAP/'03_vfx.png'),width=Inches(4.3));p(D,'Actual Unreal capture: three fountains at the three markers, after material Undo. Discrete captured state; not a live recording.',size=9)
p(D,'Preset review: one lightweight GPU sprite emitter, 90 particles/second, lifetime 1.4–1.75 seconds, sprite size 6–12 cm, default Niagara sprite material. Three new effects maximum; scale 0.25–1; anchors at least 100 cm apart. Effect Type: 5000 cm distance, maximum 4 instances per system and 8 overall. This classroom review does not certify packaged-game performance.',size=10)
page(D,'Live Unreal run sheet   4 of 4   recovery','Recover and keep teaching','Use one short recovery attempt, then switch to the prepared browser walkthrough. State what is observed and what is still unverified. Do not spend the peer-review time debugging the editor.')
table(D,['Symptom','Recovery','Teaching fallback'],[['Preview blocked','Read the message. Check target + textures separately, slot 0, valid fresh MI_ name, or three tagged anchors. Preview again.','Show the captured preview and ask students to identify scope.'],['Selection changed after Preview','Restore intended selection and Preview again. Apply consumes the old plan even when blocked.','Discuss why the old plan should expire.'],['Duplicate instance or occupied anchor','Use a new MI_ name; for effects, Undo the known previous batch. Do not overwrite assets or pile on effects.','Show the repeat-placement block capture.'],['Session expired or panel failed','Native: Reconnect, then Preview. PySide: close and reopen using w7.show_assistants().','Use the other UI only if already prepared; otherwise browser.'],['Unexpected partial result','Stop. Inspect target and protected cube, output asset and placed Actors. Failed material setup restores assignments but may retain a draft. Do not Save All.','Describe the uncertainty; use verified captured states.'],['Undo state is unclear','Do not issue repeated Undo blindly. Reopen the saved W07_MaterialVFXLab and decline saving only the rehearsal map edits. Inspect any new instance separately.','Continue from baseline capture.'],['Editor or dependency unavailable','Switch to live-demo.html. No engine, Qt, network or installation is needed.','Use prediction, observed result, repeat block and recovery prompts.']],[2.0,4.8,3.3])
matrix(D,[('Reset after class','Do not save the changed demo map. Reopen its saved baseline and verify two plain cubes and zero effects. New instance cleanup is separate: inspect references before deleting only your rehearsal asset. Never delete the Generated folder as a reset.'),('Verified in this rehearsal','UE 5.7.4 material creation/readback, protected cube/shared mesh unchanged, material Undo with asset retained, three exact VFX placements, duplicate block, VFX Undo, invalid/missing/stale selections, live PySide workflows and reopening. Native UI Preview/Apply is checked separately in the evidence report.'),('Scope of evidence','Actual editor results and discrete screenshots establish this fixture workflow. They do not establish every possible asset, target-machine performance, or a packaged-game result. See the rehearsal report beside the fallback captures for the recorded checks.')])
save(D,'Week 7 Live Demo Run Sheet.docx')
print('Created four reference-based DOCX handouts')
